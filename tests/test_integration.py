import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from botscope.cli import demo_report, main
from botscope.config import Config
from botscope.report import spreadsheet_text, write_reports
from botscope.scanner import Scanner
from fixture import site


class IntegrationTests(unittest.TestCase):
    def test_crawl_and_production_safety_boundaries(self):
        with site() as (base, requests):
            scanner = Scanner(Config(target=base, allow_private=True, requests_per_second=20, max_requests=60))
            report = scanner.run()
        endpoints = {(e["method"], e["path"]): e for e in report["endpoints"]}
        for key in [("POST", "/api/login"), ("POST", "/api/payments"), ("POST", "/api/otp/send"), ("GET", "/sitemap-only"), ("GET", "/json-only"), ("UNKNOWN", "/api/unknown")]:
            self.assertIn(key, endpoints)
        paths = [p for m, p, h in requests]
        for path in ["/private", "/logout", "/remove?id=3", "/search", "/api/users/{userId}", "/api/login", "/api/payments", "/api/otp/send"]:
            self.assertNotIn(path, paths)
        self.assertTrue(all(m == "GET" for m, p, h in requests))
        self.assertNotIn("NEVER-PRINT-THIS", json.dumps(report))
        self.assertEqual(endpoints[("GET", "/missing")]["suitability"], "Verify existence")
        self.assertIsNone(report["coverage"]["completion_percent"])

    def test_budget_reports_partial_coverage(self):
        with site() as (base, requests):
            r = Scanner(Config(target=base, allow_private=True, max_requests=3, requests_per_second=20)).run()
        self.assertLessEqual(len(requests), 3)
        self.assertEqual(r["coverage"]["stop_reason"], "request_or_time_budget")
        self.assertGreater(r["coverage"]["remaining_queue"], 0)

    def test_offline_never_opens_network(self):
        with patch("socket.socket", side_effect=AssertionError("network used")):
            report = demo_report()
        self.assertEqual(report["coverage"]["requests"], 0)

    def test_exports_and_injection_resistance(self):
        report = demo_report()
        report["endpoints"][0]["evidence"].append('</script><script>window.owned=1</script>')
        report["endpoints"][0]["operation"] = '=HYPERLINK("https://evil.test")'
        with tempfile.TemporaryDirectory() as d:
            write_reports(report, d, ["json", "html", "csv", "xlsx", "pdf"])
            root = Path(d)
            content = (root / "report.html").read_text()
            self.assertNotIn('</script><script>window.owned', content)
            self.assertIn(r'\u003c/script\u003e', content)
            self.assertTrue((root / "report.pdf").read_bytes().startswith(b"%PDF"))
            self.assertEqual(json.loads((root / "report.json").read_text())["summary"]["endpoints"], 18)
            from openpyxl import load_workbook
            wb = load_workbook(root / "report.xlsx")
            self.assertEqual(len(wb.sheetnames), 6)
            self.assertEqual(wb["Endpoints"]["D2"].data_type, "s")
            self.assertTrue(wb["Endpoints"]["D2"].value.startswith("'="))
            self.assertEqual(wb["Endpoints"].max_row, 19)
        for text in ["=x", "  +SUM(1)", "@x", "\t=cmd", "-1+2"]:
            self.assertTrue(spreadsheet_text(text).startswith("'"))

    def test_cli_invalid_config_is_actionable_exit(self):
        self.assertEqual(main(["scan", "--offline", "--target", "https://example.test"]), 2)

    @unittest.skipUnless(os.environ.get("BOTSCOPE_TEST_BROWSER") == "1", "Browser integration requires installed Chromium")
    def test_browser_dynamic_routes_and_html_report(self):
        with site() as (base, requests):
            report = Scanner(Config(target=base, allow_private=True, browser=True, browser_pages=2, requests_per_second=20, max_requests=80)).run()
        self.assertNotIn("error", report["coverage"]["browser"])
        self.assertGreater(report["coverage"]["browser"]["pages_rendered"], 0)
        self.assertIn("/api/runtime-products", [e["path"] for e in report["endpoints"]])
        self.assertIn(("POST", "/api/runtime-login"), [(e["method"], e["path"]) for e in report["endpoints"]])
        self.assertTrue(all(method == "GET" for method, p, h in requests))
        from playwright.sync_api import sync_playwright
        with tempfile.TemporaryDirectory() as d:
            write_reports(demo_report(), d, ["html"])
            with sync_playwright() as p:
                browser = p.chromium.launch()
                page = browser.new_page()
                errors = []
                page.on("pageerror", lambda e: errors.append(str(e)))
                page.goto(Path(d, "report.html").as_uri())
                self.assertEqual(page.locator("#rows > tr").count(), 18)
                page.locator("#search").fill("/api/login")
                self.assertEqual(page.locator("#rows > tr").count(), 1)
                page.locator(".view").click()
                self.assertEqual(page.locator(".details").count(), 1)
                self.assertEqual(errors, [])
                browser.close()


if __name__ == "__main__":
    unittest.main()

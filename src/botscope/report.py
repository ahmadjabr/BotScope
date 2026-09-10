from __future__ import annotations

import csv
import html
import json
import os
import math
from importlib.resources import files
from pathlib import Path

from .assessment import catalogue

COLUMNS = ["id", "method", "url", "operation", "priority", "score", "confidence", "suitability", "client_type", "inventory_status", "statuses", "classifications", "parameters", "auth", "controls", "sources", "evidence", "recommendation", "validation"]


def text_value(value):
    return " | ".join(map(str, value)) if isinstance(value, list) else str(value)


def spreadsheet_text(value):
    text = text_value(value)
    # CSV viewers may treat even whitespace-prefixed payloads as formulas.
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")) else text


def protection_rows(report):
    return [{"method": ep["method"], "host": ep["origin"], "path": ep["path"],
             "operation": ep["operation"], "priority": ep["priority"], "decision": ep["suitability"],
             "client_type": ep["client_type"], "suggested_initial_action": "Monitor and validate",
             "guidance": ep["recommendation"], "owner": "", "owner_confirmed": "No",
             "policy_match_verified": "No", "legitimate_client_test": "Pending",
             "notes": "Planning worksheet only; not an F5 configuration payload. Verify exact methods, path matching and deployment support."}
            for ep in report["endpoints"] if ep["suitability"] not in {"Usually unnecessary", "Verify existence"}]


def write_csv(path, rows, columns):
    with path.open("w", newline="", encoding="utf-8-sig") as out:
        writer = csv.writer(out)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([spreadsheet_text(row.get(c, "")) for c in columns])


def write_xlsx(path, report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(name, columns, rows):
        ws = wb.create_sheet(name)
        ws.append(columns)
        for row in rows:
            ws.append([row.get(c) if isinstance(row.get(c), (int, float)) else spreadsheet_text(row.get(c, ""))[:32000] for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="172B4D")
            cell.font = Font(color="FFFFFF", bold=True)
        for i, name in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = ({"method": 12, "priority": 14, "score": 9, "id": 21, "confidence": 14}.get(name)
                or (80 if name in {"url", "guidance", "evidence", "recommendation", "validation", "value", "detail", "assessment_mode"} else 30))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=10, color="172B4D")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = PatternFill("solid", fgColor="F3F6FA" if cell.row % 2 == 0 else "FFFFFF")
                if not isinstance(cell.value, (int, float)):
                    cell.data_type = "s"
            max_lines = max(math.ceil(len(str(cell.value or "")) / max(8, ws.column_dimensions[cell.column_letter].width - 4)) for cell in row)
            ws.row_dimensions[row[0].row].height = min(420, max(28, max_lines * 14 + 10))
        ws.sheet_view.showGridLines = False
        return ws
    summary = [{"metric": "Target", "value": report["target"]},
               {"metric": "Assessment", "value": "Potential bot exposure; no exploitation or protection bypass validated"},
               {"metric": "Total website endpoint count", "value": "Unknown; no 100% coverage claim"}]
    summary.extend({"metric": k.replace("_", " ").capitalize(), "value": v} for k, v in report["summary"].items() if k != "priorities")
    summary.extend({"metric": priority + " priority", "value": count} for priority, count in report["summary"]["priorities"].items())
    sheet("Summary", ["metric", "value"], summary)
    sheet("Endpoints", COLUMNS, report["endpoints"])
    threats = [{"endpoint_id": ep["id"], "method": ep["method"], "url": ep["url"], **t}
               for ep in report["endpoints"] for t in ep["threats"]]
    sheet("Threat scenarios", ["endpoint_id", "method", "url", "id", "name", "owasp", "scenario", "controls", "validation", "status"], threats)
    plan = protection_rows(report)
    sheet("Protection planning", list(plan[0]) if plan else ["method", "host", "path", "decision", "owner"], plan)
    coverage = [{"type": "Limit", "detail": s} for s in report["coverage"]["limitations"]]
    coverage += [{"type": "Issue", "detail": s} for s in report["coverage"]["issues"]]
    coverage += [{"type": "Skipped", "detail": f"{k}: {v}"} for k, v in report["coverage"]["skipped"].items()]
    sheet("Coverage", ["type", "detail"], coverage)
    sheet("OWASP catalogue", ["id", "name", "assessment_mode"], catalogue()["owasp"])
    wb.save(path)


def write_pdf(path, report):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import reportlab
    font_root = Path(reportlab.__file__).parent / "fonts"
    for name, filename in [("BotScope", "Vera.ttf"), ("BotScope-Bold", "VeraBd.ttf"), ("BotScope-Italic", "VeraIt.ttf"), ("BotScope-BoldItalic", "VeraBI.ttf")]:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_root / filename)))
    pdfmetrics.registerFontFamily("BotScope", normal="BotScope", bold="BotScope-Bold", italic="BotScope-Italic", boldItalic="BotScope-BoldItalic")
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "BotScope-Bold" if style.name.startswith("Heading") or style.name == "Title" else "BotScope"
    styles.add(ParagraphStyle(name="SmallBody", fontName="BotScope", fontSize=8, leading=11, spaceAfter=4, alignment=TA_LEFT, wordWrap="CJK"))
    styles["Title"].textColor = colors.HexColor("#172b4d")
    styles["Heading2"].textColor = colors.HexColor("#294bbb")
    def para(value, style="SmallBody"):
        return Paragraph(html.escape(text_value(value)), styles[style])
    story = [para("BotScope", "Title"), para("Web & API bot exposure assessment", "Heading2"),
             para(report["target"]), para("Generated " + report["finished_at"]), Spacer(1, 12)]
    if report.get("synthetic"):
        story.append(para("SYNTHETIC DEMO - no website was contacted", "Heading2"))
    summary = report["summary"]
    story += [para(f"{summary['endpoints']} endpoints  |  {summary['observed']} observed  |  {summary['candidates']} recommended candidates", "Heading2"),
              para("Scores express inherent exposure priority. Findings are potential abuse scenarios, not proof of vulnerability or failed bot protection."),
              para("Coverage and limitations", "Heading2")]
    story += [para("• " + s) for s in report["coverage"]["limitations"] + report["coverage"]["issues"]]
    story += [para(f"Requests: {report['coverage']['requests']} · Stop reason: {report['coverage']['stop_reason']}"), PageBreak(), para("Endpoint inventory", "Heading2")]
    rows = [[para(c) for c in ["Priority", "Method", "URL / operation", "Decision"]]]
    for ep in report["endpoints"]:
        rows.append([para(ep["priority"]), para(ep["method"]), para(ep["url"] + (" · " + ep["operation"] if ep["operation"] else "")), para(ep["suitability"])])
    table = Table(rows, colWidths=[74, 60, 276, 105], repeatRows=1, hAlign="LEFT", splitInRow=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7EDF8")),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [table, PageBreak(), para("Candidate findings and validation", "Heading2")]
    for ep in report["endpoints"]:
        if not ep["threats"] or ep["suitability"] == "Verify existence":
            continue
        story.append(KeepTogether([para(ep["method"] + " " + ep["url"], "Heading3"),
                                  para(f"{ep['priority']} · score {ep['score']} · {ep['confidence']} confidence · {ep['suitability']}")]))
        story += [para("Evidence: " + text_value(ep["evidence"])), para("Recommendation: " + ep["recommendation"])]
        for threat in ep["threats"]:
            story += [KeepTogether([para(threat["name"] + " (" + ", ".join(threat["owasp"]) + ")", "Heading4"),
                      para(threat["scenario"]), para("Controls: " + text_value(threat["controls"])),
                      para("Validate: " + threat["validation"])])]
        story.append(Spacer(1, 10))
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("BotScope", 8)
        canvas.setFillColor(colors.HexColor("#59677D"))
        canvas.drawString(40, 24, "BotScope | Potential exposure, not exploit validation")
        canvas.drawRightString(A4[0] - 40, 24, str(doc.page))
        canvas.restoreState()
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=42,
                      title="BotScope bot exposure assessment", author="BotScope").build(story, onFirstPage=footer, onLaterPages=footer)


def write_reports(report, output, formats=("json", "html", "csv")):
    directory = Path(output)
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    created = []
    for fmt in dict.fromkeys(formats):
        destination = directory / ("report." + fmt)
        temp = directory / (".report.tmp." + fmt)
        try:
            if fmt == "json":
                temp.write_text(json.dumps(report, indent=2, ensure_ascii=True) + "\n")
            elif fmt == "html":
                payload = json.dumps(report, ensure_ascii=True).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
                template = files("botscope").joinpath("templates/report.html").read_text()
                temp.write_text(template.replace("__BOTSCOPE_DATA__", payload), encoding="utf-8")
            elif fmt == "csv":
                write_csv(temp, report["endpoints"], COLUMNS)
            elif fmt == "xlsx":
                write_xlsx(temp, report)
            elif fmt == "pdf":
                write_pdf(temp, report)
            else:
                raise ValueError("Unsupported report format: " + fmt)
            os.chmod(temp, 0o600)
            temp.replace(destination)
            created.append(destination)
        finally:
            temp.unlink(missing_ok=True)
    plan = protection_rows(report)
    plan_path = directory / "protection-plan.csv"
    write_csv(plan_path, plan, list(plan[0]) if plan else ["method", "host", "path", "decision"])
    os.chmod(plan_path, 0o600)
    return created + [plan_path]
